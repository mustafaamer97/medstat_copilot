from docx import Document


def create_word_report(
    dataset_name,
    table1,
    results_text,
    statistical_methods
):

    document = Document()

    document.add_heading(
        "MedStat Copilot Statistical Report",
        level=1
    )

    document.add_paragraph(
        f"Dataset: {dataset_name}"
    )

    document.add_heading(
        "Statistical Methods",
        level=2
    )

    document.add_paragraph(
        statistical_methods
    )

    document.add_heading(
        "Table 1",
        level=2
    )

    rows = table1.reset_index()

    table = document.add_table(
        rows=1,
        cols=len(rows.columns)
    )

    for i, column in enumerate(rows.columns):
        table.rows[0].cells[i].text = str(column)

    for _, row in rows.iterrows():

        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = str(value)

    document.add_heading(
        "Results",
        level=2
    )

    document.add_paragraph(
        results_text
    )

    return document
